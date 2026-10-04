"""
Agent 4: Controlled Transformation and Output.

Activates only after every review item has a human decision. It:
1. builds the 17 target columns from approved column mappings;
2. applies approved cell changes (ApprovedChange from Agent 3), using
   the same normalizers that produced the reviewer's preview;
3. casts each column to its data-dictionary type, leaving missing
   values blank (never zero or "N/A");
4. writes Cleaned_SOV.xlsx (sheet "Cleaned_SOV", headers in row 1, no
   merged cells) with an Audit_Log sheet, plus Audit_Log.json;
5. validates the written file against the schema (NFR-5).

The source data table is never modified.
"""

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

from app.agents.data_quality import normalizers
from app.agents.data_quality.config import SOV_REQUIRED_FIELDS
from app.agents.data_quality.normalizers import Operation
from app.agents.data_quality.recommendation_schema import ApprovedChange


OUTPUT_FILE = "Cleaned_SOV.xlsx"
OUTPUT_SHEET = "Cleaned_SOV"
AUDIT_SHEET = "Audit_Log"
AUDIT_JSON = "Audit_Log.json"

# Data dictionary (problem statement, section 08).
FIELD_TYPES = {
    "Reference": "string",
    "Address": "string",
    "City": "string",
    "State": "string",
    "Zip": "integer",
    "County": "string",
    "Country": "string",
    "Building Value": "float",
    "Contents": "float",
    "BI": "float",
    "Occupancy": "string",
    "Construction": "string",
    "Storeys": "integer",
    "Number of Buildings": "integer",
    "Year Built": "integer",
    "Fire Sprinklers (Y/N)": "string",
    "Other": "float",
}


ALLOWED_VALUES = {
    "Fire Sprinklers (Y/N)": {"Y", "N", "Y13", "Y(13R)"},
}


class ExportBlockedError(RuntimeError):
    """Raised when export is attempted before every item is reviewed."""


class AuditEntry(BaseModel):
    source_column: Optional[str] = None
    target_column: Optional[str] = None
    source_row: Optional[int] = None
    transformation_applied: str
    before_value: Any = None
    after_value: Any = None
    confidence: Optional[float] = None
    approved_by: str
    timestamp: str
    rationale: str
    recommendation_id: Optional[str] = None


class TransformationResult(BaseModel):
    output_path: str
    audit_json_path: str
    rows: int
    applied_changes: int
    audit_entries: int
    blanked_uncastable_cells: int
    schema_valid: bool
    schema_problems: list[str] = Field(default_factory=list)


class TransformationAgent:
    def transform(
        self,
        data_table: pd.DataFrame,
        mappings: list[dict],
        changes: list[ApprovedChange],
        ledger: Any,
        require_complete: bool = True,
    ) -> tuple[pd.DataFrame, list[AuditEntry], int]:
        """
        Build the cleaned DataFrame and its audit entries.

        data_table: Agent 1's table (index = Excel row numbers).
        mappings:   Agent 2 mappings (used only through approved
                    column-mapping changes).
        changes:    ReviewSession.approved_changes().
        ledger:     ReviewLedger; export requires ledger.export_ready.
        """

        # require_complete=False builds a preview with the decisions made
        # so far; it is never written as Cleaned_SOV.xlsx.
        if require_complete and not ledger.export_ready:
            pending = ledger.pending()
            raise ExportBlockedError(
                f"{len(pending)} recommendation(s) still need a review "
                "decision; export is blocked (FR-5)."
            )

        now = datetime.now(timezone.utc).isoformat()
        audit: list[AuditEntry] = []

        # 1. Approved column mappings -----------------------------------
        target_to_source: dict[str, str] = {}

        for change in changes:
            if change.transformation_applied != Operation.RENAME_COLUMN:
                continue

            target_to_source[change.target_column] = change.source_column
            audit.append(
                AuditEntry(
                    source_column=change.source_column,
                    target_column=change.target_column,
                    transformation_applied="rename_column",
                    before_value=change.source_column,
                    after_value=change.target_column,
                    confidence=change.confidence,
                    approved_by=change.approved_by,
                    timestamp=change.timestamp.isoformat(),
                    rationale=change.rationale,
                    recommendation_id=change.recommendation_id,
                )
            )

        rows = list(data_table.index)
        output = pd.DataFrame(
            {
                field: (
                    data_table[target_to_source[field]].astype(object).tolist()
                    if field in target_to_source
                    else [None] * len(rows)
                )
                for field in SOV_REQUIRED_FIELDS
            },
            dtype=object,
        )

        # Agent 3 positions -> output positions are identical (same row
        # order); keep the Excel row for the audit log.
        excel_rows = {position: int(row) for position, row in enumerate(rows)}

        # 2. Approved cell changes --------------------------------------
        applied = 0

        for change in changes:
            if change.transformation_applied == Operation.RENAME_COLUMN:
                continue

            field = change.target_column

            if field not in target_to_source:
                raise ValueError(
                    f"{change.recommendation_id}: '{field}' has no approved "
                    "column mapping; re-run the analysis after changing "
                    "mappings."
                )

            if change.source_column and target_to_source[field] != change.source_column:
                raise ValueError(
                    f"{change.recommendation_id}: approved for source column "
                    f"'{change.source_column}' but {field} is now mapped from "
                    f"'{target_to_source[field]}'; re-run the analysis."
                )

            position = change.row_index
            current = output.at[position, field]

            if repr(_plain(current)) != repr(_plain(change.before_value)):
                raise ValueError(
                    f"{change.recommendation_id}: row {position} {field} is "
                    f"{current!r}, expected {change.before_value!r}."
                )

            output.at[position, field] = change.after_value
            applied += 1
            audit.append(
                AuditEntry(
                    source_column=change.source_column,
                    target_column=field,
                    source_row=change.source_row or excel_rows.get(position),
                    transformation_applied=change.transformation_applied.value,
                    before_value=_plain(change.before_value),
                    after_value=_plain(change.after_value),
                    confidence=change.confidence,
                    approved_by=change.approved_by,
                    timestamp=change.timestamp.isoformat(),
                    rationale=change.rationale,
                    recommendation_id=change.recommendation_id,
                )
            )

        # 3. Type casting -----------------------------------------------
        blanked = 0

        for field in SOV_REQUIRED_FIELDS:
            casted, failures, converted = _cast_column(
                output[field],
                FIELD_TYPES[field],
                ALLOWED_VALUES.get(field),
            )
            output[field] = casted

            if converted:
                audit.append(
                    AuditEntry(
                        source_column=target_to_source.get(field),
                        target_column=field,
                        transformation_applied=f"cast_to_{FIELD_TYPES[field]}",
                        before_value=f"{converted} cell(s)",
                        after_value=FIELD_TYPES[field],
                        confidence=1.0,
                        approved_by="data_dictionary",
                        timestamp=now,
                        rationale=(
                            "Lossless representation cast required by the "
                            "target data dictionary (e.g. 1970.0 -> 1970)."
                        ),
                    )
                )

            for position, value in failures:
                blanked += 1
                audit.append(
                    AuditEntry(
                        source_column=target_to_source.get(field),
                        target_column=field,
                        source_row=excel_rows.get(position),
                        transformation_applied="uncastable_value_left_blank",
                        before_value=_plain(value),
                        after_value=None,
                        confidence=1.0,
                        approved_by="data_dictionary",
                        timestamp=now,
                        rationale=(
                            f"The reviewed value is not a valid "
                            f"{FIELD_TYPES[field]} for {field}; the cell is left blank "
                            "(no placeholder) and the original value is "
                            "kept here."
                        ),
                    )
                )

        return output, audit, applied

    def export(
        self,
        data_table: pd.DataFrame,
        mappings: list[dict],
        changes: list[ApprovedChange],
        ledger: Any,
        output_dir: str | Path,
    ) -> TransformationResult:
        output, audit, applied = self.transform(
            data_table, mappings, changes, ledger
        )

        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        workbook_path = output_dir / OUTPUT_FILE
        audit_path = output_dir / AUDIT_JSON

        _write_workbook(output, audit, workbook_path)

        audit_path.write_text(
            json.dumps(
                [entry.model_dump() for entry in audit],
                indent=2,
                default=str,
            ),
            encoding="utf-8",
        )

        problems = validate_output(workbook_path)

        return TransformationResult(
            output_path=str(workbook_path),
            audit_json_path=str(audit_path),
            rows=int(len(output)),
            applied_changes=applied,
            audit_entries=len(audit),
            blanked_uncastable_cells=sum(
                1
                for entry in audit
                if entry.transformation_applied == "uncastable_value_left_blank"
            ),
            schema_valid=not problems,
            schema_problems=problems,
        )


def _plain(value: Any) -> Any:
    if normalizers.is_missing(value):
        return None

    if hasattr(value, "item") and not isinstance(value, str):
        try:
            return value.item()
        except (TypeError, ValueError):
            return value

    if isinstance(value, pd.Timestamp):
        return value.isoformat()

    return value


def _cast_column(
    series: pd.Series,
    field_type: str,
    allowed: set[str] | None = None,
) -> tuple[list[Any], list[tuple[int, Any]], int]:
    """Return (cast values, uncastable (position, value), converted count)."""

    values: list[Any] = []
    failures: list[tuple[int, Any]] = []
    converted = 0

    for position, value in enumerate(series.tolist()):
        if normalizers.is_missing(value):
            values.append(None)
            continue

        cast = _cast_value(value, field_type)

        if allowed is not None and cast not in allowed:
            cast = None

        if cast is None:
            failures.append((position, value))
            values.append(None)
            continue

        if type(cast) is not type(value) or cast != value:
            converted += 1

        values.append(cast)

    return values, failures, converted


def _cast_value(value: Any, field_type: str) -> Any:
    if field_type == "string":
        if normalizers.is_number(value):
            number = float(value)
            return str(int(number)) if number.is_integer() else str(value)

        return str(value)

    if field_type == "float":
        return (
            float(value)
            if normalizers.is_number(value)
            else None
        )

    if field_type == "integer":
        if isinstance(value, str):
            digits = value.strip()
            # ZIP+4 is kept as its 5-digit ZIP; the dictionary types Zip
            # as an integer.
            if len(digits) == 10 and digits[5] == "-" and digits[:5].isdigit():
                digits = digits[:5]
            return int(digits) if digits.isdigit() else None

        if normalizers.is_number(value) and float(value).is_integer():
            return int(value)

        return None

    return None


def _write_workbook(
    output: pd.DataFrame,
    audit: list[AuditEntry],
    path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = OUTPUT_SHEET
    sheet.append(list(SOV_REQUIRED_FIELDS))

    zip_column = SOV_REQUIRED_FIELDS.index("Zip") + 1

    for row in output.itertuples(index=False):
        sheet.append([None if normalizers.is_missing(v) else v for v in row])

    # Zip is stored as an integer (data dictionary) but displayed with
    # 5 digits so leading zeros stay visible (00802).
    for (cell,) in sheet.iter_rows(
        min_row=2, min_col=zip_column, max_col=zip_column
    ):
        if isinstance(cell.value, int):
            cell.number_format = "00000"

    audit_sheet = workbook.create_sheet(AUDIT_SHEET)
    columns = list(AuditEntry.model_fields)
    audit_sheet.append(columns)

    for entry in audit:
        data = entry.model_dump()
        audit_sheet.append(
            [
                value
                if value is None or isinstance(value, (int, float, str))
                else json.dumps(value, default=str)
                for value in (data[column] for column in columns)
            ]
        )

    workbook.save(path)


def validate_output(path: str | Path) -> list[str]:
    """Check the written workbook against the target schema (NFR-5)."""

    problems: list[str] = []
    workbook = load_workbook(path)

    if OUTPUT_SHEET not in workbook.sheetnames:
        return [f"sheet '{OUTPUT_SHEET}' is missing"]

    if workbook.sheetnames[0] != OUTPUT_SHEET:
        problems.append(f"first sheet is '{workbook.sheetnames[0]}'")

    sheet = workbook[OUTPUT_SHEET]
    header = [cell.value for cell in sheet[1]]

    if header != list(SOV_REQUIRED_FIELDS):
        problems.append(f"headers are {header}")

    if sheet.merged_cells.ranges:
        problems.append("sheet contains merged cells")

    expected = {
        "string": (str,),
        "float": (float, int),
        "integer": (int,),
    }

    for column_index, field in enumerate(SOV_REQUIRED_FIELDS, start=1):
        allowed = expected[FIELD_TYPES[field]]

        for (cell,) in sheet.iter_rows(
            min_row=2, min_col=column_index, max_col=column_index
        ):
            value = cell.value

            if value is None:
                continue

            if isinstance(value, bool) or not isinstance(value, allowed):
                problems.append(
                    f"{field} row {cell.row}: {value!r} is not "
                    f"{FIELD_TYPES[field]}"
                )
                break

            if field == "Fire Sprinklers (Y/N)" and value not in {
                "Y", "N", "Y13", "Y(13R)"
            }:
                problems.append(
                    f"{field} row {cell.row}: {value!r} is not an allowed value"
                )
                break

    return problems
