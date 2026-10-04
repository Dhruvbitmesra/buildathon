"""
End-to-end and integration regressions found by running the whole
pipeline on the real SOV files.
"""

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.agents.data_quality.agent import DataQualityAgent
from app.agents.data_quality.recommendation_schema import (
    ReviewAction,
    ReviewDecision,
)
from app.agents.schema_mapping.agent import SchemaMappingAgent, build_mapping_json
from app.agents.schema_mapping.decision_engine import resolve_mapping
from app.agents.schema_mapping.fuzzy_matcher import fuzzy_match_header
from app.agents.schema_mapping.header_normalizer import normalize_header
from app.agents.schema_mapping.llm_decision import LLMMappingDecision
from app.agents.schema_mapping.semantic_pipeline import run_semantic_pipeline
from app.agents.sheet_discovery.agent import (
    SheetDiscoveryAgent,
    clean_header_names,
    extract_table,
)
from app.agents.transformation.agent import (
    ExportBlockedError,
    TransformationAgent,
    validate_output,
)
from app.ingestion.csv_reader import read_csv_file
from app.pipeline import effective_mappings, run_pipeline
from app.state.sov_state import SOVState


SAMPLE_CSV = Path("data/sample/sample_sov.csv")
K4T9 = Path("data/input/SOV_K4T9.xlsx")


# ---------------------------------------------------------------------
# Ingestion / Agent 1
# ---------------------------------------------------------------------


def test_header_names_are_unique_and_clean():
    names = clean_header_names(["Loc # ", "Item\n#", None, "Total", "Total", ""])

    assert names == ["Loc #", "Item #", "Unnamed_3", "Total", "Total (2)", "Unnamed_6"]


def test_extract_table_excludes_blank_total_and_note_rows():
    raw = pd.DataFrame(
        [
            ["Title", None, None, None, None],
            ["Loc", "Address", "City", "State", "Building"],
            [1, "1 Main St", "Boston", "MA", 100.0],
            [None, None, None, None, None],
            [2, "2 Main St", "Boston", "MA", 200.0],
            [None, "Total", None, None, 300.0],
            ["*Values include equipment", None, None, None, None],
        ]
    )
    original = raw.copy()

    table, excluded = extract_table(raw, header_row=1)

    assert list(table.columns) == ["Loc", "Address", "City", "State", "Building"]
    assert list(table.index) == [3, 5]  # 1-based Excel rows
    assert [e["reason"] for e in excluded] == ["total_row", "note_or_footer"]
    pd.testing.assert_frame_equal(raw, original)


def test_blank_unnamed_columns_are_dropped():
    raw = pd.DataFrame([[None, "Loc", "City"], [None, 1, "Boston"], [None, 2, "Salem"]])

    table, _ = extract_table(raw, header_row=0)

    assert list(table.columns) == ["Loc", "City"]


def test_csv_cells_are_typed_but_leading_zeros_survive(tmp_path):
    path = tmp_path / "s.csv"
    path.write_text("Loc,Zip,TIV,Year\nA,02108,450000,1998.5\n", encoding="utf-8")

    sheet = read_csv_file(path)[0]["CSV_Data"]

    assert sheet.iloc[1].tolist() == ["A", "02108", 450000, 1998.5]


@pytest.mark.skipif(not K4T9.exists(), reason="sample not present")
def test_best_primary_sheet_is_selected():
    from app.ingestion.loader import load_sov_file

    result = SheetDiscoveryAgent().discover(load_sov_file(str(K4T9)).sheet_data)

    assert result.selected_sheet == "Locations"
    assert result.manifest[0].classification == "Primary"
    assert all(entry.reasoning for entry in result.manifest)
    assert any("Several sheets" in w for w in result.warnings)


def test_no_sov_sheet_is_reported_on_state():
    state = SOVState(sheet_data={"Notes": pd.DataFrame([["just a note"]])})

    state = SheetDiscoveryAgent().run(state)

    assert state.errors


# ---------------------------------------------------------------------
# Agent 2
# ---------------------------------------------------------------------


def test_hash_is_kept_as_number_word():
    assert normalize_header("Loc #") == "loc number"
    assert normalize_header("# Of Stories") == "number of stories"


def test_fuzzy_requires_key_words_to_match():
    assert fuzzy_match_header("Account Name").matched_field is None
    assert fuzzy_match_header("Constrution").matched_field == "Construction"


@pytest.mark.parametrize(
    "target, samples, contradicts",
    [
        ("Number of Buildings", ["22756726", "24854236", "5880070"], True),
        ("Number of Buildings", ["1", "2", "1"], False),
        ("Occupancy", ["1.0", "1.0", "0.5"], True),
        ("Building Value", ["Animal Shelter", "Pump", "Hall"], True),
        ("Building Value", ["1", "2", "3"], True),
        ("Building Value", ["250000", "120000", "7316910"], False),
        ("Year Built", ["9999", "9999", "1989", "1974"], False),
        ("Year Built", ["40.5", "41.2", "38.9"], True),
    ],
)
def test_value_plausibility_veto(target, samples, contradicts):
    assert SchemaMappingAgent._values_contradict(target, samples) is contradicts


def test_ambiguous_building_header_follows_values():
    agent = SchemaMappingAgent()

    assert agent._disambiguate_by_values("buildings", ["2275", "24854236", "5880"])["target_field"] == "Building Value"
    assert agent._disambiguate_by_values("building", ["1", "2", "3"])["target_field"] == "Number of Buildings"
    assert agent._disambiguate_by_values("building", ["Shed", "Hall", "Pump"]) is None


def test_no_llm_client_means_no_llm_call():
    pipeline = run_semantic_pipeline(source_header="Zqx Flarb")

    decision = resolve_mapping(pipeline, llm_client=None)

    assert decision.method in {"no_llm", "semantic"}


class FailingLLM:
    def decide(self, context):
        raise RuntimeError("network down")


def test_llm_failure_keeps_other_evidence():
    agent = SchemaMappingAgent(llm_client=FailingLLM())

    result = agent.map(
        __import__(
            "app.agents.schema_mapping.agent", fromlist=["SchemaMappingInput"]
        ).SchemaMappingInput(
            source_headers=["No. Floors", "Zqx Flarb"],
            sample_values={"No. Floors": ["1", "2", "3"], "Zqx Flarb": ["a", "b", "c"]},
        )
    )
    mapped = {m.source_header: m.target_field for m in result.mappings}

    assert mapped["No. Floors"] == "Storeys"


def test_verbal_llm_confidence_is_coerced():
    decision = LLMMappingDecision(target_field="Storeys", confidence="high", reason="x")

    assert decision.confidence == 0.85


def test_mapping_method_is_reported():
    from app.agents.schema_mapping.agent import SchemaMappingInput

    result = SchemaMappingAgent().map(
        SchemaMappingInput(source_headers=["Zip", "Year Built"])
    )

    assert {m.method for m in result.mappings} <= {"exact", "domain_alias"}


def test_mapping_json_matches_problem_statement_format():
    data = build_mapping_json(
        [
            {"source_header": "Loc #", "target_field": "Reference", "score": 0.97, "method": "fuzzy"},
            {"source_header": "Col7", "target_field": None, "score": 0.31, "method": "unresolved"},
        ],
        sheet="SOV",
        header_row=2,
    )

    assert data["header_row"] == 3
    assert data["mappings"]["Col7"] == {
        "target": None, "confidence": 0.31, "method": "none", "flag": "human_review_required",
    }
    assert data["unresolved_count"] == 1
    assert data["overall_confidence"] == 0.97


# ---------------------------------------------------------------------
# Agent 4
# ---------------------------------------------------------------------


def _reviewed(approve_all: bool):
    table = pd.DataFrame(
        {
            "Loc": ["A", "B", "C"],
            "Zip": [802.0, "02108", 75219.0],
            "TIV": ["$1,200,000", 500000, "TBD"],
            "Sprk": ["YES", "MAYBE", "N"],
            "Yr": [1970.0, 2001, 1999],
        },
        index=[5, 6, 7],
    )
    state = SOVState(
        data_table=table,
        schema_mappings=[
            {"source_header": "Loc", "target_field": "Reference", "score": 1.0, "method": "exact"},
            {"source_header": "Zip", "target_field": "Zip", "score": 1.0, "method": "exact"},
            {"source_header": "TIV", "target_field": "Building Value", "score": 0.9, "method": "fuzzy"},
            {"source_header": "Sprk", "target_field": "Fire Sprinklers (Y/N)", "score": 0.9, "method": "domain_alias"},
            {"source_header": "Yr", "target_field": "Year Built", "score": 1.0, "method": "exact"},
        ],
    )
    agent = DataQualityAgent()
    state = agent.run(state, current_year=2026)
    session = agent.session_from_state(state, 2026)

    if approve_all:
        for rec in list(session.ledger.pending()):
            session.submit(
                ReviewAction(
                    recommendation_id=rec.recommendation_id,
                    decision=ReviewDecision.APPROVE,
                    reviewer="tester",
                )
            )

    return state, session


def test_export_is_blocked_until_everything_is_reviewed(tmp_path):
    state, session = _reviewed(approve_all=False)

    with pytest.raises(ExportBlockedError):
        TransformationAgent().export(
            state.data_table, state.schema_mappings, session.approved_changes(),
            session.ledger, tmp_path,
        )


def test_export_applies_approved_changes_and_casts(tmp_path):
    state, session = _reviewed(approve_all=True)
    original = state.data_table.copy()

    result = TransformationAgent().export(
        state.data_table, state.schema_mappings, session.approved_changes(),
        session.ledger, tmp_path,
    )

    assert result.schema_valid, result.schema_problems
    pd.testing.assert_frame_equal(state.data_table, original)

    sheet = load_workbook(result.output_path)["Cleaned_SOV"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    header = [c.value for c in sheet[1]]
    col = {name: header.index(name) for name in header}

    assert rows[0][col["Building Value"]] == 1200000.0
    assert rows[2][col["Building Value"]] is None  # "TBD": blank, never 0
    assert rows[0][col["Zip"]] == 802
    assert sheet.cell(row=2, column=col["Zip"] + 1).number_format == "00000"
    assert rows[0][col["Fire Sprinklers (Y/N)"]] == "Y"
    assert rows[1][col["Fire Sprinklers (Y/N)"]] is None  # unknown kept out
    assert rows[0][col["Year Built"]] == 1970
    assert rows[0][col["County"]] is None

    audit = json.loads(Path(result.audit_json_path).read_text())
    blanked = [e for e in audit if e["transformation_applied"] == "uncastable_value_left_blank"]
    assert {e["before_value"] for e in blanked} == {"TBD", "MAYBE"}
    assert all(e["approved_by"] and e["timestamp"] for e in audit)
    assert {e["source_row"] for e in audit if e["transformation_applied"] == "parse_monetary"} == {5}


def test_validate_output_detects_problems(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Wrong"
    workbook.active.append(["Reference"])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)

    assert validate_output(path) == ["sheet 'Cleaned_SOV' is missing"]


def test_validate_output_detects_wrong_headers(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Cleaned_SOV"
    workbook.active.append(["Reference", "Zip"])
    path = tmp_path / "bad.xlsx"
    workbook.save(path)

    assert any("headers" in problem for problem in validate_output(path))


# ---------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------


def test_missing_file_gives_readable_error(tmp_path):
    result = run_pipeline(str(tmp_path / "nope.xlsx"), output_dir=tmp_path, use_llm=False)

    assert not result.ok
    assert "does not exist" in result.errors[0]


def test_unsupported_file_gives_readable_error(tmp_path):
    path = tmp_path / "x.txt"
    path.write_text("hello")

    result = run_pipeline(str(path), output_dir=tmp_path, use_llm=False)

    assert not result.ok
    assert "Unsupported file type" in result.errors[0]


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="sample not present")
def test_pipeline_stops_at_review_gate_by_default(tmp_path):
    result = run_pipeline(str(SAMPLE_CSV), output_dir=tmp_path, use_llm=False)

    assert result.ok
    assert result.stage_reached == "human_review"
    assert not (tmp_path / "Cleaned_SOV.xlsx").exists()
    assert (tmp_path / "review_queue.json").exists()
    assert (tmp_path / "schema_mapping.json").exists()


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="sample not present")
def test_pipeline_end_to_end_with_approval(tmp_path):
    result = run_pipeline(
        str(SAMPLE_CSV), output_dir=tmp_path, use_llm=False, approve="all", reviewer="tester"
    )

    assert result.ok, result.errors
    assert result.stage_reached == "transformation"
    assert result.export["schema_valid"]
    assert Path(result.artifacts["cleaned_sov"]).name == "Cleaned_SOV.xlsx"
    assert Path(result.artifacts["audit_log"]).name == "Audit_Log.json"
    assert [t["agent"] for t in result.state.agent_trace] == [
        "sheet_discovery", "schema_mapping", "data_quality", "transformation",
    ]


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="sample not present")
def test_reviewer_mapping_edit_is_revalidated(tmp_path):
    first = run_pipeline(str(SAMPLE_CSV), output_dir=tmp_path, use_llm=False)
    mapping = next(
        r for r in json.loads((tmp_path / "review_queue.json").read_text())
        if r["source_column"] == "Zip" and r["action_type"] == "column_mapping"
    )
    decisions = tmp_path / "decisions.json"
    decisions.write_text(json.dumps([
        {"recommendation_id": mapping["recommendation_id"], "decision": "edit",
         "edited_value": None, "reason": "leave Zip out"},
    ]))

    result = run_pipeline(
        str(SAMPLE_CSV), output_dir=tmp_path, use_llm=False, decisions_file=str(decisions)
    )

    assert first.ok and result.ok
    assert "Zip" not in result.state.canonical_data.columns
    assert any(
        m["source_header"] == "Zip" and m["target_field"] is None
        for m in result.state.schema_mappings
    )


def test_effective_mappings_keep_undecided_proposals():
    from app.agents.data_quality.reflection import ReviewLedger

    mappings = [{"source_header": "A", "target_field": "City", "score": 0.9, "method": "fuzzy"}]

    assert effective_mappings(mappings, ReviewLedger()) == mappings


def test_groq_client_retries_on_rate_limit(monkeypatch):
    from app.agents.schema_mapping import groq_client

    monkeypatch.setattr(groq_client.time, "sleep", lambda seconds: None)

    class RateLimited(Exception):
        status_code = 429

    calls = {"n": 0}

    class Completions:
        def create(self, **request):
            calls["n"] += 1
            if calls["n"] < 3:
                raise RateLimited("Rate limit reached. Please try again in 1.2s.")
            return "ok"

    client = groq_client.GroqLLMClient.__new__(groq_client.GroqLLMClient)
    client.client = type("C", (), {"chat": type("Ch", (), {"completions": Completions()})()})()

    assert client._create_with_retry(model="m", messages=[]) == "ok"
    assert calls["n"] == 3
    assert groq_client._retry_delay(RateLimited("try again in 194.9ms"), 1, 15) == pytest.approx(0.6949)


def test_strong_header_match_skips_llm():
    agent = SchemaMappingAgent(llm_client=FailingLLM())

    assert agent._has_strong_header_match("Yr. Built")
    assert not agent._has_strong_header_match("Zqx Flarb")


# ---------------------------------------------------------------------
# Vector memory
# ---------------------------------------------------------------------


def _fake_embed(text):
    import numpy as np

    vector = np.array([text.count(c) for c in "abcdefghijklmnopqrstuvwxyz"], dtype=float)
    return vector / (np.linalg.norm(vector) or 1.0)


def test_memory_recalls_confirmed_mapping(tmp_path):
    from app.agents.schema_mapping.mapping_memory import MappingMemory

    memory = MappingMemory(tmp_path / "m.json", embed=_fake_embed)
    memory.remember("Fire Prot.", "Fire Sprinklers (Y/N)", approved=True)
    memory.save()

    reloaded = MappingMemory(tmp_path / "m.json", embed=_fake_embed)

    assert reloaded.recall("fire prot")["target"] == "Fire Sprinklers (Y/N)"
    assert reloaded.recall("Fire Prot.")["exact"] is True
    assert reloaded.recall("Zqx") is None


def test_memory_blocks_rejected_targets(tmp_path):
    from app.agents.schema_mapping.mapping_memory import MappingMemory

    memory = MappingMemory(tmp_path / "m.json", embed=_fake_embed)
    memory.remember("Account Name", "County", approved=False)

    assert memory.blocked_targets("Account Name") == {"County"}
    assert memory.recall("Account Name") is None


def test_agent2_uses_memory_as_evidence(tmp_path):
    from app.agents.schema_mapping.agent import SchemaMappingInput
    from app.agents.schema_mapping.mapping_memory import MappingMemory

    memory = MappingMemory(tmp_path / "m.json", embed=_fake_embed)
    memory.remember("Fire Prot.", "Fire Sprinklers (Y/N)", approved=True)

    result = SchemaMappingAgent(memory=memory).map(
        SchemaMappingInput(source_headers=["Fire Prot."], sample_values={"Fire Prot.": ["Y", "N", "Y"]})
    )

    assert result.mappings[0].target_field == "Fire Sprinklers (Y/N)"
    assert result.mappings[0].method == "memory"


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="sample not present")
def test_export_stores_reviewed_mappings_in_memory(tmp_path, monkeypatch):
    from app.agents.schema_mapping.mapping_memory import MappingMemory

    memory_path = tmp_path / "memory.json"
    monkeypatch.setenv("SOV_MEMORY_PATH", str(memory_path))

    result = run_pipeline(
        str(SAMPLE_CSV), output_dir=tmp_path, use_llm=False, approve="all", reviewer="tester"
    )

    assert result.ok
    stored = MappingMemory(memory_path)
    assert stored.recall("Bldg Repl Cost New")["target"] == "Building Value"
    # Only headers are stored, never row values.
    assert "450000" not in memory_path.read_text()
